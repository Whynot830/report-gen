"""Сборка docx-отчёта из payload на базе шаблона коллеги."""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
DEFAULT_BASE = ROOT.parent
BUILTIN_STYLES = ROOT / "assets" / "gost_styles.docx"

# Свободные id: в gost_styles.docx заняты num 1–94 и abstract 0–79.
HEADING_NUM_ABSTRACT = 200
HEADING_NUM_ID = 200
REF_ABSTRACT = 201
REF_NUM_ID = 201
REF_STYLE_NAME = "References List"
REF_STYLE_ID = "ReferencesList"
UNORDERED_STYLE = "Unordered List Paragraph"
FIGURE_NUM_RE = re.compile(r"^Рисунок\s+\d+(?:\.\d+)?\s*[–-]\s*")


def h1(text: str) -> dict:
    return {"t": "h1", "text": text}


def h2(text: str) -> dict:
    return {"t": "h2", "text": text}


def h3(text: str) -> dict:
    return {"t": "h3", "text": text}


def p(text: str) -> dict:
    return {"t": "p", "text": text}


def ul(*items: str) -> dict:
    return {"t": "ul", "items": list(items)}


def table(caption: str, headers: list, rows: list, continue_caption: str | None = None) -> dict:
    block = {"t": "table", "caption": caption,
             "headers": headers, "rows": rows}
    if continue_caption:
        block["continue_caption"] = continue_caption
    return block


def fig(path: str | Path, caption: str, width_cm: float = 16.5) -> dict:
    return {"t": "figure", "path": str(path), "caption": caption, "width_cm": width_cm}


def listing(caption: str, code: str) -> dict:
    return {"t": "listing", "caption": caption, "code": code}


def refs(*items: str) -> dict:
    return {"t": "refs", "items": list(items)}


def _set_run_font(run, *, size=14, bold=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _clear_first_line_override_needed(p):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    if qn("w:firstLineChars") in ind.attrib:
        del ind.attrib[qn("w:firstLineChars")]


def _clear_title_first_line(doc: Document) -> None:
    """На титуле нет красной строки: first line 0 у всех абзацев, включая таблицы."""
    for p_el in doc.element.body.iter(qn("w:p")):
        _clear_first_line_override_needed(Paragraph(p_el, doc))


def _set_paragraph_text(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        run = para.add_run(text)
        _set_run_font(run, size=14)


def _student_group(student_label: str, fallback: str = "") -> str:
    marker = "группы "
    if marker in student_label:
        return student_label.split(marker, 1)[1].strip()
    return fallback


def _fill_title_student(doc: Document, student_label: str, student_name: str) -> None:
    group = _student_group(student_label)
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            left = row.cells[0]
            right = row.cells[1]
            for para in left.paragraphs:
                if "Выполнил студент группы" in para.text:
                    _set_paragraph_text(
                        para, f"Выполнил студент группы {group}")
                    name_para = next((p for p in reversed(
                        right.paragraphs) if p.text.strip()), None)
                    if name_para is not None:
                        _set_paragraph_text(name_para, student_name)
                    return
                if "Студент" in para.text and para.runs:
                    para.runs[0].text = student_label
                    for run in para.runs[1:]:
                        run.text = ""
                    for rpara in right.paragraphs:
                        if rpara.text.strip() and rpara.runs:
                            rpara.runs[0].text = student_name
                            for run in rpara.runs[1:]:
                                run.text = ""
                            break
                    return


def _part_by_rel_suffix(document: Document, suffix: str):
    for rel in document.part.rels.values():
        if rel.reltype.endswith(suffix):
            return rel.target_part
    return None


def _apply_builtin_styles(doc: Document) -> None:
    """Подставить styles.xml и numbering.xml из вшитого ГОСТ-шаблона."""
    donor = Document(str(BUILTIN_STYLES))
    src_styles = _part_by_rel_suffix(donor, "/styles")
    dst_styles = _part_by_rel_suffix(doc, "/styles")
    if src_styles is None or dst_styles is None:
        raise RuntimeError("Не найден styles.xml у шаблона или у вшитых стилей")
    dst_styles._element = deepcopy(src_styles.element)
    doc.styles._element = dst_styles._element

    src_num = _part_by_rel_suffix(donor, "/numbering")
    dst_num = _part_by_rel_suffix(doc, "/numbering")
    if src_num is None:
        raise RuntimeError("Во вшитых стилях нет numbering.xml")
    if dst_num is None:
        raise RuntimeError("В титульном шаблоне нет numbering.xml")
    dst_num._element = deepcopy(src_num.element)


def _apply_page_setup(doc: Document) -> None:
    """Поля страницы по ГОСТ: левое 3 см, правое 1,5 см, верх/низ 2 см."""
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.gutter = Cm(0)
        section.header_distance = Cm(1)
        section.footer_distance = Cm(1)


def _clone_title(
    template: Path,
    work_no: int,
    student_label: str,
    student_name: str,
    discipline: str | None = None,
) -> Document:
    doc = Document(str(template))
    body = doc.element.body
    children = list(body)
    keep_count = 0
    for i, child in enumerate(children):
        if "Москва 2026" in "".join(child.itertext()):
            keep_count = i + 1
            break
    for child in children[keep_count:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    for para in doc.paragraphs:
        if "Практическая работа №" in para.text:
            _set_paragraph_text(para, f"Практическая работа № {work_no}")
        elif discipline and "по дисциплине" in para.text:
            _set_paragraph_text(para, f"по дисциплине «{discipline}»")

    _fill_title_student(doc, student_label, student_name)
    return doc


def _set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def _ensure_num(numbering, num_id: int, abstract_id: int) -> None:
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(num_id):
            return
    node = OxmlElement("w:num")
    node.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    node.append(abs_ref)
    numbering.append(node)


def _insert_abstract(numbering, abstract) -> None:
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)


def _add_heading_numbering(doc: Document) -> None:
    numbering = doc.part.numbering_part._element
    for absn in numbering.findall(qn("w:abstractNum")):
        if absn.get(qn("w:abstractNumId")) == str(HEADING_NUM_ABSTRACT):
            return

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(HEADING_NUM_ABSTRACT))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    levels = [
        ("0", "Heading2", "%1"),
        ("1", "Heading3", "%1.%2"),
    ]
    for ilvl, pstyle, lvl_text in levels:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), ilvl)
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), pstyle)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        pPr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:right"), "0")
        ind.set(qn("w:firstLine"), "709")
        pPr.append(ind)
        for child in (start, num_fmt, p_style, suff, text, lvl_jc, pPr):
            lvl.append(child)
        abstract.append(lvl)

    _insert_abstract(numbering, abstract)
    _ensure_num(numbering, HEADING_NUM_ID, HEADING_NUM_ABSTRACT)


def _apply_heading_base_styles(doc: Document) -> None:
    """К стилям H2/H3 из ГОСТ-шаблона добавить только автонумерацию 1 / 1.1."""
    for name, numbered_ilvl in (("Heading 2", "0"), ("Heading 3", "1")):
        style = doc.styles[name]
        pPr = style.element.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            style.element.append(pPr)
        old = pPr.find(qn("w:numPr"))
        if old is not None:
            pPr.remove(old)
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), numbered_ilvl)
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(HEADING_NUM_ID))
        numPr.append(ilvl)
        numPr.append(num_id)
        pPr.append(numPr)


def _set_style_font_size(style, half_points: str = "28") -> None:
    rPr = style.element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        style.element.append(rPr)
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), half_points)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), half_points)
    rPr.append(sz)
    rPr.append(sz_cs)


def _apply_table_name_style(doc: Document) -> None:
    try:
        _set_style_font_size(doc.styles["Table Name"], "28")
    except KeyError:
        return


def _apply_img_caption_style(doc: Document) -> None:
    try:
        _set_style_font_size(doc.styles["IMG Caption"], "28")
    except KeyError:
        return


def _add_references_style(doc: Document) -> None:
    numbering = doc.part.numbering_part._element
    exists = any(
        absn.get(qn("w:abstractNumId")) == str(REF_ABSTRACT)
        for absn in numbering.findall(qn("w:abstractNum"))
    )
    if not exists:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(REF_ABSTRACT))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), REF_STYLE_ID)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "%1.")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        pPr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:right"), "0")
        ind.set(qn("w:firstLine"), "709")
        pPr.append(ind)
        for child in (start, num_fmt, p_style, suff, lvl_text, lvl_jc, pPr):
            lvl.append(child)
        abstract.append(lvl)
        _insert_abstract(numbering, abstract)
        _ensure_num(numbering, REF_NUM_ID, REF_ABSTRACT)

    names = [s.name for s in doc.styles]
    if REF_STYLE_NAME not in names:
        style = doc.styles.add_style(REF_STYLE_NAME, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles[REF_STYLE_NAME]
    style.element.set(qn("w:customStyle"), "1")
    style.element.set(qn("w:styleId"), REF_STYLE_ID)
    style.base_style = doc.styles["Normal"]
    style.quick_style = True

    pPr = style.element.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        style.element.append(pPr)
    else:
        for child in list(pPr):
            pPr.remove(child)

    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(REF_NUM_ID))
    numPr.append(ilvl)
    numPr.append(num_id)
    pPr.append(numPr)

    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)

    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLine"), "709")
    pPr.append(ind)

    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "both")
    pPr.append(jc)


def _add_toc(doc: Document) -> None:
    title = doc.add_paragraph()
    title.style = doc.styles["Centered"]
    run = title.add_run("СОДЕРЖАНИЕ")
    _set_run_font(run, size=14, bold=True)

    p = doc.add_paragraph()
    _clear_first_line_override_needed(p)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    run_begin = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld_begin)

    run_instr = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    run_instr._r.append(instr)

    run_sep = p.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    hint = p.add_run("Правый щелчок по содержанию → Обновить поле")
    _set_run_font(hint, size=12)

    run_end = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def _add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    run = p.add_run(text)
    _set_run_font(run, size=14, bold=True)


def _add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(text)
    _set_run_font(run, size=14)


def _add_ul(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.style = doc.styles[UNORDERED_STYLE]
        run = p.add_run(item)
        _set_run_font(run, size=14)


def _add_blank(doc: Document) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    _clear_first_line_override_needed(p)


def _style_names(doc: Document) -> set[str]:
    return {s.name for s in doc.styles}


def _add_caption(doc: Document, text: str, *, figure: bool = False, listing: bool = False) -> None:
    names = _style_names(doc)
    if figure:
        text = FIGURE_NUM_RE.sub("", text)
        style_name = "IMG Caption" if "IMG Caption" in names else "Table Name"
    elif listing:
        style_name = "Listing Name" if "Listing Name" in names else "Table Name"
    else:
        style_name = "Table Name"
    p = doc.add_paragraph()
    try:
        p.style = doc.styles[style_name]
    except KeyError:
        p.style = doc.styles["Normal"]
    if figure:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _clear_first_line_override_needed(p)
    run = p.add_run(text)
    _set_run_font(run, size=14)


def _set_cell_border(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    tcPr.append(borders)


def _set_cell_margins(cell, cm: float = 0.15) -> None:
    twips = str(int(round(cm * 567)))
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), twips)
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _clear_first_line_override_needed(p)
    run = p.add_run(str(text))
    _set_run_font(run, size=12, bold=bold)
    _set_cell_border(cell)
    _set_cell_margins(cell, 0.15)


def _set_table_cell_margins(table, cm: float = 0.15) -> None:
    twips = str(int(round(cm * 567)))
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    tblCellMar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), twips)
        node.set(qn("w:type"), "dxa")
        tblCellMar.append(node)
    tblPr.append(tblCellMar)


def _add_table(doc: Document, headers: list, rows: list) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_cell_margins(table, 0.15)
    for i, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], header, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            _set_cell_text(table.rows[r_i + 1].cells[c_i], val, bold=False)


def _add_figure(doc: Document, path: str | Path, width_cm: float) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _clear_first_line_override_needed(p)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def _add_refs(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.style = doc.styles[REF_STYLE_NAME]
        run = p.add_run(item)
        _set_run_font(run, size=14)


def _add_listing(doc: Document, code: str) -> None:
    lines = code.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]
    para = doc.add_paragraph()
    names = _style_names(doc)
    if "Code" in names:
        para.style = doc.styles["Code"]
    else:
        para.style = doc.styles["Normal"]
    _clear_first_line_override_needed(para)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines or [""]):
        run = para.add_run(line if line else " ")
        if "Code" not in names:
            _set_run_font(run, size=10, name="Courier New")
        if i < len(lines) - 1:
            run.add_break()


def _render_blocks(doc: Document, blocks: list[dict], base: Path) -> None:
    for block in blocks:
        kind = block["t"]
        if kind == "h1":
            _add_heading(doc, block["text"], 1)
        elif kind == "h2":
            _add_heading(doc, block["text"], 2)
        elif kind == "h3":
            _add_heading(doc, block["text"], 3)
        elif kind == "p":
            _add_body(doc, block["text"])
        elif kind == "ul":
            _add_ul(doc, block["items"])
        elif kind == "table":
            _add_caption(doc, block["caption"])
            _add_table(doc, block["headers"], block["rows"])
            if block.get("continue_caption"):
                _add_caption(doc, block["continue_caption"])
            _add_blank(doc)
        elif kind == "figure":
            path = Path(block["path"])
            if not path.is_absolute():
                path = (base / path).resolve()
            _add_figure(doc, path, block.get("width_cm", 16.5))
            _add_caption(doc, block["caption"], figure=True)
            _add_blank(doc)
        elif kind == "listing":
            _add_caption(doc, block["caption"], listing=True)
            _add_listing(doc, block["code"])
            _add_blank(doc)
        elif kind == "refs":
            _add_refs(doc, block["items"])
        else:
            raise ValueError(f"Неизвестный тип блока: {kind}")


def build_report(payload: dict, *, base: Path | None = None) -> Path:
    """Собрать docx из payload и вернуть путь к файлу."""
    base = Path(base) if base is not None else DEFAULT_BASE
    meta = payload["meta"]
    template = Path(meta["template"])
    if not template.is_absolute():
        template = base / template
    output = Path(meta["output"])
    if not output.is_absolute():
        output = base / output

    doc = _clone_title(
        template,
        work_no=meta["work_no"],
        student_label=meta.get("student_label", ""),
        student_name=meta.get("student_name", ""),
        discipline=meta.get("discipline"),
    )
    _apply_builtin_styles(doc)
    _clear_title_first_line(doc)
    _apply_page_setup(doc)
    _set_update_fields(doc)
    _add_heading_numbering(doc)
    _apply_heading_base_styles(doc)
    _apply_table_name_style(doc)
    _apply_img_caption_style(doc)
    _add_references_style(doc)
    _add_toc(doc)
    _render_blocks(doc, payload["blocks"], base)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output
